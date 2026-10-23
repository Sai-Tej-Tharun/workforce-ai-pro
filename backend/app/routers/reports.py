"""Module 19: Reports & export (employees, attrition, payroll, attendance, leave) as XLSX / PDF / CSV."""
import csv, io
from datetime import date
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Employee, Payroll, User, ROLE_ADMIN, ROLE_HR
from ..security import staff, hr_admin
from ..services.common import audit, employee_scope_filter, latest_predictions
from .attendance import attendance_report_rows
from .leave import leave_report_rows

router = APIRouter(prefix="/reports", tags=["Reports & Export"])

MEDIA = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf", "csv": "text/csv"}


def _to_xlsx(title: str, cols: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, col in enumerate(cols, 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]]) + 3
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(width, 40)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_pdf(title: str, cols: list[str], rows: list[list]) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24, title=title)
    st = getSampleStyleSheet()
    data = [cols] + [[("" if v is None else str(v))[:28] for v in r] for r in rows[:1000]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")])]))
    doc.build([Paragraph(f"WorkForce AI Pro – {title}", st["Title"]),
               Paragraph(f"Generated {date.today().isoformat()} · {len(rows)} rows", st["Normal"]), Spacer(1, 10), t])
    return buf.getvalue()


def _to_csv(cols, rows) -> bytes:
    b = io.StringIO()
    w = csv.writer(b)
    w.writerow(cols)
    w.writerows(rows)
    return b.getvalue().encode()


def _respond(title: str, cols: list[str], rows: list[list], fmt: str, db, user, request):
    if fmt not in MEDIA:
        raise HTTPException(422, "format must be xlsx, pdf or csv")
    data = {"xlsx": _to_xlsx, "pdf": _to_pdf}.get(fmt, lambda t, c, r: _to_csv(c, r))(title, cols, rows)
    audit(db, user, "EXPORT", "report", title, {"format": fmt, "rows": len(rows)}, request)
    return StreamingResponse(io.BytesIO(data), media_type=MEDIA[fmt],
                             headers={"Content-Disposition": f'attachment; filename="{title.lower().replace(" ", "_")}_{date.today()}.{fmt}"'})


@router.get("/employees")
def employees_report(request: Request, format: str = "xlsx", department_id: Optional[int] = None, status: Optional[str] = None,
                     db: Session = Depends(get_db), user: User = Depends(staff)):
    q = select(Employee)
    scope = employee_scope_filter(user)
    if scope is not None: q = q.where(scope)
    if department_id: q = q.where(Employee.department_id == department_id)
    if status: q = q.where(Employee.status == status)
    emps = db.scalars(q.order_by(Employee.id)).unique().all()
    show_sal = user.role in (ROLE_ADMIN, ROLE_HR)
    cols = ["Code", "Name", "Email", "Department", "Role", "Hire date", "Status", "Tenure (y)"] + (["Monthly salary"] if show_sal else [])
    rows = [[e.employee_code, e.full_name, e.email, e.department.name if e.department else "", e.role.title if e.role else "",
             e.hire_date.isoformat(), e.status, e.years_at_company] + ([e.monthly_salary] if show_sal else []) for e in emps]
    return _respond("Employee Report", cols, rows, format, db, user, request)


@router.get("/attrition")
def attrition_report(request: Request, format: str = "xlsx", risk_level: Optional[str] = None,
                     db: Session = Depends(get_db), user: User = Depends(staff)):
    q = select(Employee).where(Employee.status != "terminated")
    scope = employee_scope_filter(user)
    if scope is not None: q = q.where(scope)
    emps = db.scalars(q).unique().all()
    preds = latest_predictions(db, [e.id for e in emps])
    rows = []
    for e in emps:
        p = preds.get(e.id)
        if not p or (risk_level and p.risk_level != risk_level):
            continue
        rows.append([e.employee_code, e.full_name, e.department.name if e.department else "", f"{p.risk_score * 100:.1f}%", p.risk_level.upper(),
                     f"{p.confidence * 100:.1f}%", ", ".join(f["label"] for f in (p.top_factors or [])[:3]), p.created_at.strftime("%Y-%m-%d %H:%M")])
    rows.sort(key=lambda r: -float(r[3].rstrip("%")))
    return _respond("Attrition Risk Report", ["Code", "Name", "Department", "Risk", "Level", "Confidence", "Top drivers", "Predicted at"], rows, format, db, user, request)


@router.get("/payroll")
def payroll_report(request: Request, month: int = Query(default_factory=lambda: date.today().month, ge=1, le=12),
                   year: int = Query(default_factory=lambda: date.today().year), format: str = "xlsx",
                   db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    ps = db.scalars(select(Payroll).where(Payroll.month == month, Payroll.year == year).order_by(Payroll.employee_id)).unique().all()
    rows = [[p.employee.employee_code, p.employee.full_name, p.basic, p.hra, p.allowances, p.overtime_pay, p.bonus + p.incentive,
             p.gross, p.tax, p.pf_deduction, p.other_deductions, p.net_pay, p.status] for p in ps]
    return _respond(f"Payroll Report {year}-{month:02d}", ["Code", "Name", "Basic", "HRA", "Allow.", "OT pay", "Bonus+Inc.", "Gross", "Tax", "PF", "Other ded.", "Net", "Status"], rows, format, db, user, request)


@router.get("/attendance")
def attendance_report(request: Request, month: int = Query(default_factory=lambda: date.today().month, ge=1, le=12),
                      year: int = Query(default_factory=lambda: date.today().year), format: str = "xlsx",
                      department_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(staff)):
    rows = [[r["employee_code"], r["employee_name"], r["days_present"], r["late_days"], r["absent_days"], r["total_hours"], r["overtime_hours"]]
            for r in attendance_report_rows(db, user, month, year, department_id)]
    return _respond(f"Attendance Report {year}-{month:02d}", ["Code", "Name", "Present", "Late", "Absent", "Hours", "Overtime h"], rows, format, db, user, request)


@router.get("/leave")
def leave_report(request: Request, year: int = Query(default_factory=lambda: date.today().year), format: str = "xlsx",
                 db: Session = Depends(get_db), user: User = Depends(hr_admin)):
    rows = [[r["employee_code"], r["employee_name"], r["leave_type"], r["days"], r["requests"]] for r in leave_report_rows(db, year)]
    return _respond(f"Leave Report {year}", ["Code", "Name", "Leave type", "Days", "Requests"], rows, format, db, user, request)
